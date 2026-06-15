from __future__ import annotations

import argparse
import os
import sys
import time
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from ugocr.api import create_app
from ugocr.config import Settings


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run local API smoke tests without starting uvicorn.")
    parser.add_argument("--config", type=Path, default=Path("configs/ugocr.example.yaml"))
    parser.add_argument("--chinese-image", type=Path, default=Path("wtest2.png"))
    parser.add_argument("--table-image", type=Path, default=Path("table.png"))
    parser.add_argument("--use-vlm", action="store_true", help="Call the configured VLM during OCR smoke tests.")
    parser.add_argument("--only", choices=["all", "ocr", "table", "translate"], default="all")
    parser.add_argument("--app-key", default="test_key")
    parser.add_argument("--max-seconds", type=float, default=5.0)
    return parser.parse_args()


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    args = parse_args()
    os.environ["UGOCR_CONFIG"] = str(args.config)
    settings = Settings.from_env()
    results: list[tuple[str, bool, str]] = []

    startup = time.perf_counter()
    with TestClient(create_app(settings)) as client:
        startup_elapsed = time.perf_counter() - startup
        print(f"INFO startup_warmup_sec={startup_elapsed:.3f}", flush=True)

        if args.only in {"all", "ocr"}:
            _record_tuple(
                results,
                _post_image(
                    client,
                    name="POST /api/v1/ocr/recognize",
                    url="/api/v1/ocr/recognize",
                    image_path=args.chinese_image,
                    use_vlm=args.use_vlm,
                    app_key=args.app_key,
                    max_seconds=args.max_seconds,
                    validator=lambda data: data.get("target_lang") == "ch" and isinstance(data.get("result"), list),
                ),
            )
        if args.only in {"all", "table"}:
            _record_tuple(results, _post_table_excel(client, args.table_image, args.use_vlm, args.app_key, args.max_seconds))
        if args.only == "translate":
            _record_tuple(results, _post_translate(client, args.app_key, args.max_seconds))

    if not all(ok for _, ok, _ in results):
        raise SystemExit(1)


def _record(results: list[tuple[str, bool, str]], name: str, ok: bool, detail: str) -> None:
    results.append((name, ok, detail))
    mark = "OK" if ok else "FAIL"
    print(f"{mark:4} {name:<22} {detail}", flush=True)


def _record_tuple(results: list[tuple[str, bool, str]], item: tuple[str, bool, str]) -> None:
    _record(results, item[0], item[1], item[2])


def _post_image(
    client: TestClient,
    *,
    name: str,
    url: str,
    image_path: Path,
    use_vlm: bool,
    app_key: str,
    max_seconds: float,
    validator,
) -> tuple[str, bool, str]:
    path = _project_path(image_path)
    if not path.exists():
        return name, False, f"sample image missing: {path}"
    with path.open("rb") as handle:
        started = time.perf_counter()
        response = client.post(
            url,
            headers={"appKey": app_key},
            files={"file": (path.name, handle, "image/png")},
            data={"source_lang": "ch", "use_vlm": "true" if use_vlm else "false"},
        )
        elapsed = time.perf_counter() - started
    if response.status_code != 200:
        return name, False, _brief_response(response)
    data = response.json()
    ok = bool(validator(data)) and elapsed <= max_seconds
    detail = f"elapsed={elapsed:.3f}s lines={len(data.get('result', []))}"
    return name, ok, detail


def _post_table_excel(client: TestClient, image_path: Path, use_vlm: bool, app_key: str, max_seconds: float) -> tuple[str, bool, str]:
    name = "POST /api/v1/ocr/table"
    path = _project_path(image_path)
    if not path.exists():
        return name, False, f"sample image missing: {path}"
    with path.open("rb") as handle:
        started = time.perf_counter()
        response = client.post(
            "/api/v1/ocr/table",
            headers={"appKey": app_key},
            files={"file": (path.name, handle, "image/png")},
            data={"source_lang": "ch", "use_vlm": "true" if use_vlm else "false"},
        )
        elapsed = time.perf_counter() - started
    if response.status_code != 200:
        return name, False, _brief_response(response)
    output = PROJECT_ROOT / "data" / "outputs" / "smoke-table.xlsx"
    output.write_bytes(response.content)
    try:
        workbook = load_workbook(output)
        sheet = workbook.active
        detail = f"elapsed={elapsed:.3f}s xlsx_opened=true rows={sheet.max_row} cols={sheet.max_column} sheet={sheet.title!r}"
        return name, elapsed <= max_seconds, detail
    except Exception as exc:
        return name, False, f"xlsx_opened=false error={exc}"


def _post_translate(client: TestClient, app_key: str, max_seconds: float) -> tuple[str, bool, str]:
    name = "POST /api/v1/translate/text"
    started = time.perf_counter()
    response = client.post(
        "/api/v1/translate/text",
        headers={"appKey": app_key},
        json={"texts": ["سلام"], "target_lang": "ch"},
    )
    elapsed = time.perf_counter() - started
    if response.status_code != 200:
        return name, False, _brief_response(response)
    data = response.json()
    ok = (
        data.get("target_lang") == "ch"
        and isinstance(data.get("texts"), list)
        and isinstance(data.get("translate_results"), list)
        and len(data.get("texts", [])) == len(data.get("translate_results", []))
        and elapsed <= max_seconds
    )
    return name, ok, f"elapsed={elapsed:.3f}s items={len(data.get('translate_results', []))} source={data.get('source_lang')}"


def _brief_response(response: Any) -> str:
    text = response.text.strip().replace("\n", " ")
    if len(text) > 500:
        text = text[:497] + "..."
    return f"status={response.status_code} body={text}"


def _project_path(path: Path) -> Path:
    if path.is_absolute():
        return path
    return (PROJECT_ROOT / path).resolve()


if __name__ == "__main__":
    main()
