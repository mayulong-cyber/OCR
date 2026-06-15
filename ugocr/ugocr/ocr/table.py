from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image

from ugocr.config import Settings
from ugocr.dependencies import import_required, require_path
from ugocr.ocr.paddle_engine import (
    PaddleOCRTextEngine,
    _paddlex_device,
    _paddlex_predictor_option,
    configure_paddle_runtime_env,
    paddlex_ocr_config,
)
from ugocr.ocr.postprocess import VLMPostProcessor
from ugocr.ocr.types import InternalTableCell
from ugocr.utils.images import image_to_numpy, load_image, sort_text_boxes


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[dict[str, Any]]] = []
        self._current_row: list[dict[str, Any]] | None = None
        self._current_cell: dict[str, Any] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        name = tag.lower()
        if name == "tr":
            self._current_row = []
            return
        if name in {"td", "th"} and self._current_row is not None:
            attr_map = {key.lower(): value for key, value in attrs if value is not None}
            self._current_cell = {
                "text": "",
                "rowspan": _positive_int(attr_map.get("rowspan"), 1),
                "colspan": _positive_int(attr_map.get("colspan"), 1),
                "header": name == "th",
            }

    def handle_data(self, data: str) -> None:
        if self._current_cell is not None:
            self._current_cell["text"] += data

    def handle_endtag(self, tag: str) -> None:
        name = tag.lower()
        if name in {"td", "th"} and self._current_cell is not None and self._current_row is not None:
            self._current_cell["text"] = " ".join(str(self._current_cell["text"]).split())
            self._current_row.append(self._current_cell)
            self._current_cell = None
            return
        if name == "tr" and self._current_row is not None:
            self.rows.append(self._current_row)
            self._current_row = None


class PaddlexTableEngine:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self._engine: Any | None = None

    def _load(self) -> Any:
        if self._engine is not None:
            return self._engine
        require_path(self.settings.paddle.det_model_dir, "PP-OCRv5 detection model directory")
        require_path(self.settings.paddle.table_rec_model_dir, "table OCR recognition model directory")
        require_path(self.settings.paddle.table_model_dir, "SLANet table model directory")
        configure_paddle_runtime_env(self.settings.runtime)
        paddlex_inference = import_required("paddlex.inference", "requirements-gpu.txt or requirements-cpu.txt")
        create_pipeline = getattr(paddlex_inference, "create_pipeline")
        self._engine = create_pipeline(
            config=_paddlex_table_config(self.settings),
            device=None if _paddlex_device(self.settings.runtime) == "cpu" else _paddlex_device(self.settings.runtime),
            pp_option=_paddlex_predictor_option(paddlex_inference, self.settings.runtime),
        )
        return self._engine

    def predict_html(self, image: Image.Image) -> str | None:
        engine = self._load()
        result = list(
            engine.predict(
                image_to_numpy(image),
                use_doc_orientation_classify=False,
                use_doc_unwarping=False,
                use_layout_detection=False,
                use_ocr_model=True,
                text_det_limit_side_len=self.settings.paddle.det_limit_side_len,
                text_det_limit_type="max",
                text_rec_score_thresh=0.0,
            )
        )
        return _find_html(result)

    def warmup(self) -> None:
        self._load()


class TablePipeline:
    def __init__(self, settings: Settings, postprocessor: VLMPostProcessor) -> None:
        self.settings = settings
        self.structure_engine = PaddlexTableEngine(settings)
        self.ocr_engine = PaddleOCRTextEngine(
            runtime=settings.runtime,
            paddle=settings.paddle,
            rec_model_dir=settings.paddle.table_rec_model_dir,
            char_dict_path=settings.paddle.table_char_dict_path,
            engine_name="ppocrv5-table-rec",
        )
        self.postprocessor = postprocessor

    def warmup(self) -> None:
        self.structure_engine.warmup()
        self.ocr_engine.warmup()

    def recognize_to_excel(self, image_bytes: bytes, use_vlm: bool | None = None) -> tuple[Path, list[InternalTableCell], bool, list[str]]:
        payload = load_image(image_bytes)
        warnings: list[str] = []
        cells: list[InternalTableCell]
        try:
            html = self.structure_engine.predict_html(payload.image)
            if html:
                cells = cells_from_html(html)
            else:
                warnings.append("SLANet did not return HTML, using OCR grid reconstruction.")
                cells = self._grid_from_ocr(payload.image)
        except Exception as exc:
            warnings.append(f"SLANet table parsing failed: {exc}. Using OCR grid reconstruction.")
            cells = self._grid_from_ocr(payload.image)
        if use_vlm is not True:
            corrected = False
            warnings.append("VLM postprocess skipped by request.")
        else:
            vlm_result = self.postprocessor.enhance_table(image_bytes, cells)
            cells = vlm_result.cells
            corrected = vlm_result.corrected
            warnings.extend(vlm_result.warnings)
        target = self.settings.runtime.output_dir / f"table-{uuid4().hex}.xlsx"
        write_xlsx(cells, target)
        return target, cells, corrected, warnings

    def _grid_from_ocr(self, image: Image.Image) -> list[InternalTableCell]:
        lines = self.ocr_engine.recognize(image)
        records = sort_text_boxes(
            [{"text": line.text, "score": line.score, "box": line.box, "engine": line.engine} for line in lines],
        )
        if len(records) == 0:
            return []
        row_groups: list[list[dict[str, Any]]] = []
        heights = []
        for item in records:
            ys = [p[1] for p in item["box"]]
            heights.append(max(1.0, max(ys) - min(ys)))
        tolerance = max(8.0, sorted(heights)[len(heights) // 2] * 0.65)
        for item in records:
            ys = [p[1] for p in item["box"]]
            cy = sum(ys) / len(ys)
            placed = False
            for row in row_groups:
                row_y = sum(sum(p[1] for p in entry["box"]) / len(entry["box"]) for entry in row) / len(row)
                if abs(cy - row_y) <= tolerance:
                    row.append(item)
                    placed = True
                    break
            if not placed:
                row_groups.append([item])
        cells: list[InternalTableCell] = []
        for row_idx, row in enumerate(row_groups, start=1):
            row.sort(key=lambda item: min(p[0] for p in item["box"]))
            for col_idx, item in enumerate(row, start=1):
                cells.append(
                    InternalTableCell(
                        row=row_idx,
                        col=col_idx,
                        text=str(item["text"]).strip(),
                        score=float(item["score"]),
                        box=item["box"],
                    )
                )
        return cells


def cells_from_html(html: str) -> list[InternalTableCell]:
    parser = _HTMLTableParser()
    parser.feed(html)
    occupied: set[tuple[int, int]] = set()
    cells: list[InternalTableCell] = []
    for row_idx, row in enumerate(parser.rows, start=1):
        col_idx = 1
        for raw in row:
            while (row_idx, col_idx) in occupied:
                col_idx += 1
            rowspan = int(raw["rowspan"])
            colspan = int(raw["colspan"])
            cell = InternalTableCell(
                row=row_idx,
                col=col_idx,
                text=str(raw["text"]).strip(),
                score=1.0,
                rowspan=rowspan,
                colspan=colspan,
            )
            cells.append(cell)
            for row_offset in range(rowspan):
                for col_offset in range(colspan):
                    occupied.add((row_idx + row_offset, col_idx + col_offset))
            col_idx += colspan
    return cells


def write_xlsx(cells: list[InternalTableCell], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "OCR Table"
    thin = Side(style="thin", color="D9DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type="solid", fgColor="EEF2FF")
    for cell in cells:
        excel_cell = sheet.cell(row=cell.row, column=cell.col, value=cell.text)
        excel_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        excel_cell.border = border
        if cell.row == 1:
            excel_cell.font = Font(bold=True)
            excel_cell.fill = header_fill
        if cell.rowspan > 1 or cell.colspan > 1:
            end_row = cell.row + cell.rowspan - 1
            end_col = cell.col + cell.colspan - 1
            sheet.merge_cells(start_row=cell.row, start_column=cell.col, end_row=end_row, end_column=end_col)
    max_row = max((cell.row + cell.rowspan - 1 for cell in cells), default=1)
    max_col = max((cell.col + cell.colspan - 1 for cell in cells), default=1)
    for row in range(1, max_row + 1):
        sheet.row_dimensions[row].height = 24
    for col in range(1, max_col + 1):
        width = 12
        for cell in cells:
            if cell.col == col:
                width = max(width, min(32, len(cell.text) * 1.7 + 4))
        sheet.column_dimensions[sheet.cell(row=1, column=col).column_letter].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_diagnostic_xlsx(path: Path, title: str, messages: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "UGOCR Error"
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, color="9C1C1C")
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FEE2E2")
    sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    for index, message in enumerate(messages, start=3):
        sheet.cell(row=index, column=1, value=index - 2)
        cell = sheet.cell(row=index, column=2, value=message)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=index, start_column=2, end_row=index, end_column=4)
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 90
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 16
    for row in range(1, max(4, len(messages) + 4)):
        sheet.row_dimensions[row].height = 28
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _find_html(result: Any) -> str | None:
    if result is None:
        return None
    if not isinstance(result, (str, bytes, dict, list, tuple)):
        try:
            html_value = getattr(result, "html")
        except Exception:
            html_value = None
        if html_value is not None:
            found = _find_html(html_value)
            if found is not None:
                return found
        try:
            json_value = getattr(result, "json")
        except Exception:
            json_value = None
        if json_value is not None:
            found = _find_html(json_value)
            if found is not None:
                return found
    if isinstance(result, str) and "<table" in result.lower():
        return result
    if isinstance(result, dict):
        for key in ("html", "html_str", "table_html"):
            value = result.get(key)
            if isinstance(value, str) and "<table" in value.lower():
                return value
            if isinstance(value, list):
                joined = "".join(str(item) for item in value)
                if "<table" in joined.lower():
                    return joined
        for key in ("res", "result", "data"):
            value = result.get(key)
            found = _find_html(value)
            if found is not None:
                return found
    if isinstance(result, (list, tuple)):
        for item in result:
            found = _find_html(item)
            if found is not None:
                return found
    return None


def _paddlex_table_config(settings: Settings) -> dict[str, Any]:
    return {
        "pipeline_name": "table_recognition",
        "use_doc_preprocessor": False,
        "use_layout_detection": False,
        "use_ocr_model": True,
        "SubModules": {
            "TableStructureRecognition": {
                "module_name": "table_structure_recognition",
                "model_name": "SLANet",
                "model_dir": str(settings.paddle.table_model_dir),
            }
        },
        "SubPipelines": {
            "GeneralOCR": paddlex_ocr_config(settings.paddle, settings.paddle.table_rec_model_dir)
        },
    }


def _positive_int(value: str | None, default: int) -> int:
    try:
        parsed = int(value) if value is not None else default
    except ValueError:
        parsed = default
    return max(1, parsed)
